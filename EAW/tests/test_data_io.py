"""Excel 导入/导出测试。"""

from datetime import date
from io import BytesIO

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from EAW.models import Category, Item, Proficiency

from .factories import create_item, create_user_with_defaults


def _xlsx_bytes(rows):
    """rows[0] 为表头,构造一个 Excel 文件并返回字节。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(client, rows, filename='import.xlsx', **extra):
    content = _xlsx_bytes(rows)
    upload = SimpleUploadedFile(filename, content,
                                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    return client.post(reverse('import_user_data'), {'file': upload, **extra})


class ExportTests(TestCase):
    def setUp(self):
        self.user = create_user_with_defaults('alice')
        self.client.login(username='alice', password='pass-12345678')

    def test_export_contains_only_own_items(self):
        other = create_user_with_defaults('bob')
        create_item(self.user, item='apple', content='苹果')
        create_item(other, item='bobs-secret')

        response = self.client.get(reverse('export_user_data'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

        workbook = openpyxl.load_workbook(BytesIO(response.content))
        sheet_rows = list(workbook.active.iter_rows(values_only=True))
        items = [row[0] for row in sheet_rows[1:]]
        self.assertIn('apple', items)
        self.assertNotIn('bobs-secret', items)

    def test_export_headers(self):
        response = self.client.get(reverse('export_user_data'))
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        headers = [cell.value for cell in workbook.active[1]]
        self.assertEqual(headers[0], 'Item')
        self.assertIn('US Phonetic', headers)


class ImportTests(TestCase):
    def setUp(self):
        self.user = create_user_with_defaults('alice')
        self.client.login(username='alice', password='pass-12345678')

    def test_import_creates_items(self):
        rows = [
            ['Item', 'Content', 'Input Date', 'Init Date', 'Proficiency', 'Category'],
            ['apple', '苹果', date(2025, 2, 1), date(2025, 2, 1), 'Unfamiliar', '单词'],
            ['run', '跑', date(2025, 2, 1), date(2025, 2, 1), 'Mastered', '单词'],
        ]
        response = _upload(self.client, rows)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Item.objects.filter(user=self.user).count(), 2)
        run = Item.objects.get(user=self.user, item='run')
        self.assertEqual(run.proficiency, Proficiency.MASTERED)

    def test_import_creates_missing_category(self):
        rows = [
            ['Item', 'Category'],
            ['定语从句', '语法'],
        ]
        _upload(self.client, rows)
        self.assertTrue(Category.objects.filter(user=self.user, name='语法').exists())
        item = Item.objects.get(user=self.user, item='定语从句')
        self.assertEqual(item.category.name, '语法')

    def test_import_skips_rows_without_item(self):
        rows = [
            ['Item', 'Content'],
            ['', 'no name'],
            ['apple', '苹果'],
        ]
        _upload(self.client, rows)
        self.assertEqual(Item.objects.filter(user=self.user).count(), 1)

    def test_import_missing_required_column(self):
        rows = [
            ['NotItem', 'Content'],
            ['apple', '苹果'],
        ]
        response = _upload(self.client, rows)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Item.objects.count(), 0)
        self.assertContains(response, '缺少必要的列')

    def test_import_without_file_shows_error(self):
        response = self.client.post(reverse('import_user_data'), {})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '请求无效')

    def test_import_does_not_fetch_definitions_by_default(self):
        rows = [
            ['Item', 'Content', 'Category'],
            ['apple', '苹果', '单词'],
        ]
        with self.settings():
            # fetch_definitions 未勾选时不应触发网络调用——直接上传播 mock 以暴露意外调用
            from unittest.mock import patch
            with patch('EAW.views.data_io.fetch_and_merge_translation',
                       side_effect=AssertionError('should not fetch')):
                _upload(self.client, rows)
        self.assertEqual(Item.objects.filter(user=self.user).count(), 1)

    def test_roundtrip_export_then_import(self):
        """导出的文件应能无损导回(迁移/备份场景)。"""
        create_item(self.user, item='apple', content='苹果\n水果',
                    input_date=date(2025, 3, 1), init_date=date(2025, 3, 1))
        export_response = self.client.get(reverse('export_user_data'))

        # 导入前清空数据,模拟数据恢复
        Item.objects.filter(user=self.user).delete()
        upload = SimpleUploadedFile(
            'user_data.xlsx', export_response.content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.client.post(reverse('import_user_data'), {'file': upload})

        item = Item.objects.get(user=self.user, item='apple')
        self.assertEqual(item.content, '苹果\n水果')
        self.assertEqual(item.inputDate, date(2025, 3, 1))
        self.assertEqual(item.category.name, '单词')
