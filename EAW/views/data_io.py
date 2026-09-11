"""Excel 导入 / 导出视图。"""

import openpyxl
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import render
from django.utils.timezone import now

from ..models import Category, Item, Proficiency
from ..utils import fetch_and_merge_translation


@login_required
def export_user_data_to_excel(request):
    # 获取当前用户数据
    user = request.user
    items = Item.objects.filter(user=user)

    # 创建一个新的 Excel 工作簿
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "User Data"

    # 写入表头
    headers = ["Item", "Content", "Input Date", "Init Date", "Proficiency", "Category", "TTS URL", "US Phonetic", "UK Phonetic"]
    sheet.append(headers)

    # 写入用户数据
    for item in items:
        sheet.append([
            item.item,
            item.content,
            item.inputDate,
            item.initDate,
            item.get_proficiency_display(),
            item.category.name if item.category else "",
            item.src_tts,
            item.us_phonetic,
            item.uk_phonetic,
        ])

    # 创建 HTTP 响应
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="user_data.xlsx"'

    # 将工作簿保存到响应中
    workbook.save(response)

    return response


@login_required
def import_items_from_excel(request):
    if request.method == "POST" and request.FILES.get("file"):
        file = request.FILES["file"]
        user = request.user

        # 检查是否选择了“获取释义”选项
        fetch_definitions = "fetch_definitions" in request.POST

        try:
            workbook = openpyxl.load_workbook(file)
            sheet = workbook.active
        except Exception as e:
            messages.error(request, f"文件读取失败: {str(e)}")
            return render(request, "import_data.html")

        headers = [cell.value for cell in sheet[1]]
        required_columns = ["Item"]
        if not all(col in headers for col in required_columns):
            messages.error(request, "文件格式错误，缺少必要的列。")
            return render(request, "import_data.html")

        column_index = {header: headers.index(header) for header in headers}
        items_to_create = []
        errors = []

        # Proficiency 字段的映射
        proficiency_map = {
            "Unfamiliar": Proficiency.UNFAMILIAR,
            "Mastered": Proficiency.MASTERED,
        }

        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            try:
                item_name = row[column_index["Item"]]
                if not item_name:
                    errors.append(f"第 {row_idx} 行缺少 Item 字段，已跳过。")
                    continue

                content_index = column_index.get("Content")
                content = row[content_index] if content_index is not None else ""
                # 替换 _x000D_ 字符为换行符，先检查是否为 None
                if content:
                    content = content.replace("_x000D_", "\n")
                else:
                    content = ""  # 如果 content 为 None，赋空字符串

                input_date_index = column_index.get("Input Date")
                input_date = row[input_date_index] if input_date_index is not None else now().date()

                init_date_index = column_index.get("Init Date")
                init_date = row[init_date_index] if init_date_index is not None else now().date()

                # 处理 Proficiency 字段
                proficiency_index = column_index.get("Proficiency")
                proficiency_name = row[proficiency_index] if proficiency_index is not None else None
                proficiency_degree = proficiency_map.get(proficiency_name, Proficiency.UNFAMILIAR)

                category_index = column_index.get("Category")
                category_name = row[category_index] if category_index is not None else ""

                # 获取分类对象
                category = None
                if category_name:
                    categories = Category.objects.filter(name=category_name, user=user)
                    if categories.exists():
                        category = categories.first()
                    else:
                        category = Category.objects.create(name=category_name, user=user)

                # 初始化这些字段为默认值
                src_tts = ""
                phonetic_am = ""
                phonetic_en = ""

                # 只有类别为“单词”的条目才调用获取翻译的功能
                if fetch_definitions and category and category.name == "单词":
                    updated_content, src_tts, phonetic_am, phonetic_en = fetch_and_merge_translation(item_name, content)
                    # 更新内容
                    content = updated_content

                item = Item(
                    user=user,
                    item=item_name,
                    content=content,
                    inputDate=input_date,
                    initDate=init_date,
                    proficiency=proficiency_degree,
                    category=category,
                    src_tts=src_tts,
                    us_phonetic=phonetic_am,
                    uk_phonetic=phonetic_en,
                )
                items_to_create.append(item)

            except Exception as e:
                errors.append(f"第 {row_idx} 行处理失败: {str(e)}")

        try:
            with transaction.atomic():
                Item.objects.bulk_create(items_to_create)
            success_message = f"导入完成。成功导入 {len(items_to_create)} 条记录，{len(errors)} 条记录跳过。"
            messages.success(request, success_message)
        except Exception as e:
            messages.error(request, f"保存失败: {str(e)}")

        return render(request, "import_data.html", {
            "import_results": {
                "success_count": len(items_to_create),
                "errors": errors,
            }
        })

    messages.error(request, "请求无效，请上传文件。")
    return render(request, "import_data.html")
